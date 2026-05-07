# =============================================================================
# Face2Shop Analytics Engine — Module 1: Automated Report Generator
# Pulls from MySQL views → generates branded Excel report with charts
# =============================================================================

import os
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import seaborn as sns
from sqlalchemy import create_engine, text
import mysql.connector
from dotenv import load_dotenv
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
import warnings
warnings.filterwarnings('ignore')

# -----------------------------------------------------------------------------
# CONFIG
# -----------------------------------------------------------------------------
load_dotenv()
DB_PASS = os.getenv("DB_PASS")

COLORS = {
    'navy':      '#1B2B4B',
    'terracotta':'#C8956C',
    'teal':      '#1E4D5C',
    'coral':     '#E8503A',
    'cream':     '#FAF6F0',
    'card':      '#FFFFFF',
    'text':      '#6B6B6B',
}

OUTPUT_DIR = "outputs/reports"
CHART_DIR  = "outputs/reports/charts"
os.makedirs(CHART_DIR, exist_ok=True)

REPORT_DATE = datetime.now().strftime("%Y-%m-%d")
REPORT_FILE = f"{OUTPUT_DIR}/face2shop_report_{REPORT_DATE}.xlsx"

# -----------------------------------------------------------------------------
# DATABASE CONNECTION
# -----------------------------------------------------------------------------
def get_engine():
    return create_engine(
        "mysql+mysqlconnector://",
        creator=lambda: mysql.connector.connect(
            user="root",
            password=DB_PASS,
            host="127.0.0.1",
            port=3306,
            database="face2shop_db"
        )
    )

# -----------------------------------------------------------------------------
# DATA EXTRACTION
# -----------------------------------------------------------------------------
def extract_data(engine):
    print("  Extracting data from MySQL views...")
    queries = {
        'funnel':    "SELECT * FROM vw_conversion_funnel",
        'room':      "SELECT * FROM vw_room_performance",
        'price_seg': "SELECT * FROM vw_price_segment_analysis",
        'brand':     "SELECT * FROM vw_brand_performance",
        'daily':     "SELECT * FROM vw_daily_trend",
        'hourly':    "SELECT * FROM vw_hourly_behavior",
        'gap':       "SELECT * FROM vw_high_view_no_purchase",
    }
    data = {}
    with engine.connect() as conn:
        for key, sql in queries.items():
            data[key] = pd.read_sql(text(sql), conn)
            print(f"    ✓ {key}: {len(data[key])} rows")
    return data

# -----------------------------------------------------------------------------
# KPI SUMMARY
# -----------------------------------------------------------------------------
def extract_kpis(engine):
    print("  Extracting KPIs from fact table...")
    sql = """
        SELECT
            SUM(revenue)                                      AS total_revenue,
            COUNT(*)                                          AS total_events,
            SUM(CASE WHEN event_type='purchase' THEN 1 END)  AS total_purchases,
            SUM(CASE WHEN event_type='view'     THEN 1 END)  AS total_views,
            SUM(CASE WHEN event_type='cart'     THEN 1 END)  AS total_cart_adds,
            COUNT(DISTINCT user_id)                           AS unique_users,
            AVG(CASE WHEN event_type='purchase'
                     THEN revenue END)                        AS avg_order_value
        FROM fact_user_events
    """
    with engine.connect() as conn:
        kpis = pd.read_sql(text(sql), conn).iloc[0]
    conversion_rate = (kpis['total_purchases'] / kpis['total_views'] * 100
                       if kpis['total_views'] else 0)
    cart_abandon = ((kpis['total_cart_adds'] - kpis['total_purchases']) /
                    kpis['total_cart_adds'] * 100
                    if kpis['total_cart_adds'] else 0)
    kpis['conversion_rate']      = round(conversion_rate, 2)
    kpis['cart_abandonment_rate'] = round(cart_abandon, 2)
    return kpis

# -----------------------------------------------------------------------------
# CHART GENERATORS
# -----------------------------------------------------------------------------
def hex_to_rgb(hex_color):
    h = hex_color.lstrip('#')
    return tuple(int(h[i:i+2], 16)/255 for i in (0, 2, 4))

def save_chart(fig, name):
    path = f"{CHART_DIR}/{name}.png"
    fig.savefig(path, dpi=150, bbox_inches='tight',
                facecolor=COLORS['cream'], edgecolor='none')
    plt.close(fig)
    return path

def chart_funnel(data):
    df = data['funnel'].copy()
    # Ensure correct funnel order
    order = {'view': 0, 'cart': 1, 'purchase': 2}
    if 'event_type' in df.columns:
        df['_order'] = df['event_type'].map(order)
        df = df.sort_values('_order')
        labels = df['event_type'].str.capitalize().tolist()
        values = df['event_count'].tolist()
    else:
        labels = df.iloc[:, 0].tolist()
        values = df.iloc[:, 1].tolist()

    fig, ax = plt.subplots(figsize=(7, 4), facecolor=COLORS['cream'])
    bar_colors = [COLORS['navy'], COLORS['teal'], COLORS['terracotta']]
    bars = ax.barh(labels[::-1], values[::-1], color=bar_colors,
                   height=0.5, edgecolor='none')
    for bar, val in zip(bars, values[::-1]):
        ax.text(bar.get_width() * 1.01, bar.get_y() + bar.get_height()/2,
                f'{int(val):,}', va='center', fontsize=10,
                color=COLORS['navy'], fontweight='bold')
    ax.set_facecolor(COLORS['cream'])
    ax.spines[['top','right','left','bottom']].set_visible(False)
    ax.tick_params(left=False, bottom=False)
    ax.set_xticklabels([])
    ax.set_title('Conversion Funnel', fontsize=13, fontweight='bold',
                 color=COLORS['navy'], pad=12)
    fig.tight_layout()
    return save_chart(fig, 'funnel')

def chart_room(data):
    df = data['room'].copy()
    df = df.groupby('room_type', as_index=False)['conversion_rate_pct'].mean()
    df = df.sort_values('conversion_rate_pct', ascending=False).head(8)
    room_col = 'room_type'
    cvr_col  = 'conversion_rate_pct' 

    fig, ax = plt.subplots(figsize=(7, 4), facecolor=COLORS['cream'])
    bars = ax.bar(df[room_col], df[cvr_col],
                  color=COLORS['teal'], edgecolor='none', width=0.6)
    bars[0].set_color(COLORS['terracotta'])  # highlight top
    ax.set_facecolor(COLORS['cream'])
    ax.spines[['top','right']].set_visible(False)
    ax.set_title('Room Category Performance (CVR %)', fontsize=13,
                 fontweight='bold', color=COLORS['navy'], pad=12)
    ax.set_ylabel('Conversion Rate %', color=COLORS['text'], fontsize=9)
    ax.tick_params(axis='x', rotation=30, labelsize=8)
    ax.tick_params(axis='y', labelsize=8)
    fig.tight_layout()
    return save_chart(fig, 'room')

def chart_price_segment(data):
    df = data['price_seg'].copy()
    seg_col = df.columns[0]
    rev_col = [c for c in df.columns if 'rev' in c.lower()]
    rev_col = rev_col[0] if rev_col else df.columns[1]
    cvr_col = [c for c in df.columns if 'conv' in c.lower() or 'rate' in c.lower()]
    cvr_col = cvr_col[0] if cvr_col else df.columns[2]

    fig, ax1 = plt.subplots(figsize=(7, 4), facecolor=COLORS['cream'])
    ax2 = ax1.twinx()
    x = range(len(df))
    ax1.bar(x, df[rev_col], color=COLORS['navy'], alpha=0.8,
            edgecolor='none', width=0.4, label='Revenue')
    ax2.plot(x, df[cvr_col], color=COLORS['coral'], marker='o',
             linewidth=2, markersize=6, label='CVR %')
    ax1.set_xticks(list(x))
    ax1.set_xticklabels(df[seg_col], fontsize=8, rotation=15)
    ax1.set_facecolor(COLORS['cream'])
    ax1.spines[['top']].set_visible(False)
    ax2.spines[['top']].set_visible(False)
    ax1.set_ylabel('Revenue', color=COLORS['navy'], fontsize=9)
    ax2.set_ylabel('CVR %', color=COLORS['coral'], fontsize=9)
    ax1.set_title('Price Segment: Revenue vs Conversion', fontsize=13,
                  fontweight='bold', color=COLORS['navy'], pad=12)
    fig.tight_layout()
    return save_chart(fig, 'price_segment')

def chart_hourly(data):
    df = data['hourly'].copy()
    hr_col  = df.columns[0]
    val_col = [c for c in df.columns if 'purch' in c.lower() or 'count' in c.lower()]
    val_col = val_col[0] if val_col else df.columns[1]

    fig, ax = plt.subplots(figsize=(7, 4), facecolor=COLORS['cream'])
    ax.fill_between(df[hr_col], df[val_col],
                    color=COLORS['teal'], alpha=0.3)
    ax.plot(df[hr_col], df[val_col],
            color=COLORS['teal'], linewidth=2)
    peak_idx = df[val_col].idxmax()
    ax.scatter(df.loc[peak_idx, hr_col], df.loc[peak_idx, val_col],
               color=COLORS['coral'], zorder=5, s=80)
    ax.annotate(f"Peak: {df.loc[peak_idx, hr_col]}h",
                xy=(df.loc[peak_idx, hr_col], df.loc[peak_idx, val_col]),
                xytext=(10, 10), textcoords='offset points',
                color=COLORS['coral'], fontsize=9, fontweight='bold')
    ax.set_facecolor(COLORS['cream'])
    ax.spines[['top','right']].set_visible(False)
    ax.set_title('Hourly Purchase Behavior', fontsize=13,
                 fontweight='bold', color=COLORS['navy'], pad=12)
    ax.set_xlabel('Hour of Day', color=COLORS['text'], fontsize=9)
    ax.set_ylabel('Purchases', color=COLORS['text'], fontsize=9)
    fig.tight_layout()
    return save_chart(fig, 'hourly')

def chart_daily(data):
    df = data['daily'].copy()
    date_col = df.columns[0]
    rev_col  = [c for c in df.columns if 'rev' in c.lower()]
    rev_col  = rev_col[0] if rev_col else df.columns[1]
    df[date_col] = pd.to_datetime(df[date_col])

    fig, ax = plt.subplots(figsize=(7, 4), facecolor=COLORS['cream'])
    ax.plot(df[date_col], df[rev_col],
            color=COLORS['navy'], linewidth=2)
    ax.fill_between(df[date_col], df[rev_col],
                    color=COLORS['navy'], alpha=0.1)
    ax.set_facecolor(COLORS['cream'])
    ax.spines[['top','right']].set_visible(False)
    ax.set_title('Daily Revenue Trend', fontsize=13,
                 fontweight='bold', color=COLORS['navy'], pad=12)
    ax.set_ylabel('Revenue', color=COLORS['text'], fontsize=9)
    plt.xticks(rotation=30, fontsize=7)
    fig.tight_layout()
    return save_chart(fig, 'daily')

# -----------------------------------------------------------------------------
# EXCEL WRITER
# -----------------------------------------------------------------------------
def style_header_row(ws, row, cols, bg_hex, font_hex='FFFFFF', font_size=11):
    fill = PatternFill("solid", fgColor=bg_hex.lstrip('#'))
    font = Font(bold=True, color=font_hex.lstrip('#'), size=font_size,
                name='Georgia')
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal='center', vertical='center')

def style_data_row(ws, row, cols, bg_hex='FFFFFF', font_hex='6B6B6B'):
    fill = PatternFill("solid", fgColor=bg_hex.lstrip('#'))
    font = Font(color=font_hex.lstrip('#'), size=10, name='Calibri')
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal='center', vertical='center')

def add_chart_to_sheet(ws, img_path, anchor):
    img = XLImage(img_path)
    img.width  = 480
    img.height = 270
    ws.add_image(img, anchor)

def write_kpi_sheet(wb, kpis):
    ws = wb.create_sheet("KPI Summary")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 22

    # Title
    ws.merge_cells('A1:B1')
    title_cell = ws['A1']
    title_cell.value = "Face2Shop — KPI Summary"
    title_cell.font  = Font(bold=True, size=16, color='1B2B4B', name='Georgia')
    title_cell.fill  = PatternFill("solid", fgColor='FAF6F0')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 36

    ws.merge_cells('A2:B2')
    ws['A2'].value = f"Generated: {REPORT_DATE}"
    ws['A2'].font  = Font(size=9, color='6B6B6B', name='Calibri')
    ws['A2'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 18

    kpi_items = [
        ("Total Revenue",         f"${kpis['total_revenue']:,.0f}"),
        ("Total Views",           f"{int(kpis['total_views']):,}"),
        ("Total Purchases",       f"{int(kpis['total_purchases']):,}"),
        ("Total Cart Adds",       f"{int(kpis['total_cart_adds']):,}"),
        ("Unique Users",          f"{int(kpis['unique_users']):,}"),
        ("Conversion Rate",       f"{kpis['conversion_rate']}%"),
        ("Cart Abandonment Rate", f"{kpis['cart_abandonment_rate']}%"),
        ("Avg Order Value",       f"${kpis['avg_order_value']:,.2f}"),
    ]

    nav   = PatternFill("solid", fgColor='1B2B4B')
    cream = PatternFill("solid", fgColor='FAF6F0')
    card  = PatternFill("solid", fgColor='FFFFFF')
    thin  = Border(bottom=Side(style='thin', color='E0E0E0'))

    style_header_row(ws, 3, 2, '#1B2B4B', 'FFFFFF', 12)
    ws['A3'].value = "Metric"
    ws['B3'].value = "Value"
    ws.row_dimensions[3].height = 24

    for i, (metric, value) in enumerate(kpi_items):
        r = i + 4
        ws.cell(row=r, column=1).value = metric
        ws.cell(row=r, column=2).value = value
        bg = 'FAF6F0' if i % 2 == 0 else 'FFFFFF'
        style_data_row(ws, r, 2, f'#{bg}')
        ws.cell(row=r, column=1).font = Font(bold=True, size=10,
                                              color='1B2B4B', name='Calibri')
        ws.row_dimensions[r].height = 22

def write_data_sheet(wb, df, sheet_name, title):
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    ws.merge_cells(f'A1:{get_column_letter(len(df.columns))}1')
    ws['A1'].value = title
    ws['A1'].font  = Font(bold=True, size=14, color='1B2B4B', name='Georgia')
    ws['A1'].fill  = PatternFill("solid", fgColor='FAF6F0')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    style_header_row(ws, 2, len(df.columns), '#1B2B4B')
    for col_idx, col_name in enumerate(df.columns, 1):
        ws.cell(row=2, column=col_idx).value = str(col_name).replace('_', ' ').title()
        ws.column_dimensions[get_column_letter(col_idx)].width = 20
    ws.row_dimensions[2].height = 22

    for row_idx, row in enumerate(df.itertuples(index=False), 3):
        bg = 'FAF6F0' if row_idx % 2 == 0 else 'FFFFFF'
        style_data_row(ws, row_idx, len(df.columns), f'#{bg}')
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = value
            if isinstance(value, float):
                cell.number_format = '#,##0.00'
            elif isinstance(value, int):
                cell.number_format = '#,##0'
        ws.row_dimensions[row_idx].height = 18

def write_charts_sheet(wb, chart_paths):
    ws = wb.create_sheet("Visual Dashboard")
    ws.sheet_view.showGridLines = False

    # Title
    ws.merge_cells('A1:L1')
    ws['A1'].value = "Face2Shop — Visual Dashboard"
    ws['A1'].font  = Font(bold=True, size=16, color='1B2B4B', name='Georgia')
    ws['A1'].fill  = PatternFill("solid", fgColor='FAF6F0')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 36

    positions = ['B3', 'H3', 'B22', 'H22', 'B41']
    for path, anchor in zip(chart_paths, positions):
        if os.path.exists(path):
            add_chart_to_sheet(ws, path, anchor)

    # Set row/col sizes for chart layout
    for col in 'ABCDEFGHIJKL':
        ws.column_dimensions[col].width = 11
    for r in range(1, 60):
        ws.row_dimensions[r].height = 15

def write_gap_sheet(wb, data):
    df = data['gap'].copy()
    ws = wb.create_sheet("Gap Analysis")
    ws.sheet_view.showGridLines = False

    ws.merge_cells(f'A1:{get_column_letter(len(df.columns))}1')
    ws['A1'].value = "High View / Zero Purchase Products — Action Required"
    ws['A1'].font  = Font(bold=True, size=14, color='E8503A', name='Georgia')
    ws['A1'].fill  = PatternFill("solid", fgColor='FAF6F0')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    style_header_row(ws, 2, len(df.columns), '#C8956C', 'FFFFFF')
    for col_idx, col_name in enumerate(df.columns, 1):
        ws.cell(row=2, column=col_idx).value = str(col_name).replace('_', ' ').title()
        ws.column_dimensions[get_column_letter(col_idx)].width = 22
    ws.row_dimensions[2].height = 22

    for row_idx, row in enumerate(df.itertuples(index=False), 3):
        style_data_row(ws, row_idx, len(df.columns), 'FAF6F0' if row_idx % 2 == 0 else 'FFFFFF')
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = value
            # Highlight coral if views > 1000
            if col_idx == 2 and isinstance(value, (int, float)) and value > 1000:
                cell.fill = PatternFill("solid", fgColor='E8503A')
                cell.font = Font(bold=True, color='FFFFFF', size=10)
        ws.row_dimensions[row_idx].height = 18

# -----------------------------------------------------------------------------
# MAIN
# -----------------------------------------------------------------------------
def main():
    print("\n" + "="*60)
    print("  Face2Shop Report Generator — Starting")
    print("="*60)

    engine = get_engine()

    print("\n[1/4] Extracting data...")
    data = extract_data(engine)
    kpis = extract_kpis(engine)

    print("\n[2/4] Generating charts...")
    chart_paths = [
        chart_funnel(data),
        chart_room(data),
        chart_price_segment(data),
        chart_hourly(data),
        chart_daily(data),
    ]
    print(f"    ✓ {len(chart_paths)} charts saved to {CHART_DIR}")

    print("\n[3/4] Building Excel report...")
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    write_kpi_sheet(wb, kpis)
    write_charts_sheet(wb, chart_paths)
    write_data_sheet(wb, data['funnel'],    "Funnel Data",        "Conversion Funnel")
    write_data_sheet(wb, data['room'],      "Room Performance",   "Room Category Performance")
    write_data_sheet(wb, data['price_seg'], "Price Segments",     "Price Segment Analysis")
    write_data_sheet(wb, data['brand'],     "Brand Performance",  "Brand Performance")
    write_data_sheet(wb, data['hourly'],    "Hourly Behavior",    "Hourly Purchase Behavior")
    write_gap_sheet(wb, data)

    wb.save(REPORT_FILE)
    print(f"    ✓ Report saved: {REPORT_FILE}")

    print("\n[4/4] Summary")
    print(f"    Revenue:         ${kpis['total_revenue']:,.0f}")
    print(f"    Views:           {int(kpis['total_views']):,}")
    print(f"    Purchases:       {int(kpis['total_purchases']):,}")
    print(f"    Conversion Rate: {kpis['conversion_rate']}%")
    print(f"    Avg Order Value: ${kpis['avg_order_value']:,.2f}")
    print("\n" + "="*60)
    print("  Report generation COMPLETE")
    print("="*60 + "\n")

if __name__ == "__main__":
    main()
